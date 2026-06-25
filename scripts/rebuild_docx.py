"""使用 python-docx 从 md 源文件重建三个 .docx 文档，修复格式问题。

解决的问题：
- 所有段落均为 Normal 样式，标题没有 Heading 层级
- 编号重复（如两个"3."）
- 标题和正文无法区分
"""

import os
import re
import subprocess
import tempfile
from docx import Document
from docx.shared import Pt, Cm, Emu, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm

# Try to use SimHei for Chinese in matplotlib
for fname in fm.findSystemFonts():
    if "SimHei" in fname or "simhei" in fname:
        plt.rcParams["font.sans-serif"] = ["SimHei"]
        break
plt.rcParams["axes.unicode_minus"] = False

DOCS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs")
MERMAID_IMG_DIR = os.path.join(DOCS_DIR, "_mermaid_img")
CHART_IMG_DIR = os.path.join(DOCS_DIR, "_chart_img")

# Font/size constants
TITLE_FONT = "SimHei"
TITLE_SIZE = Pt(36)
SUBTITLE_SIZE = Pt(22)
INFO_FONT = "SimSun"
INFO_SIZE = Pt(15)
HEADING1_SIZE = Pt(16)  # 三号
HEADING2_SIZE = Pt(14)  # 四号
HEADING3_SIZE = Pt(12)  # 小四
BODY_SIZE = Pt(12)       # 小四
TABLE_HEADER_SIZE = Pt(12)
TABLE_CELL_SIZE = Pt(10.5)  # 五号
CODE_FONT = "Consolas"


def render_mermaid_to_png(mermaid_text, output_path):
    """用 mmdc 将 mermaid 文本渲染为 PNG 图片。"""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with tempfile.NamedTemporaryFile(mode="w", suffix=".mmd", encoding="utf-8", delete=False) as tmp:
        tmp.write(mermaid_text)
        tmp_path = tmp.name
    try:
        # On Windows, mmdc is a .cmd wrapper; use shell=True to resolve it
        result = subprocess.run(
            f'mmdc -i "{tmp_path}" -o "{output_path}" -b white --scale 2',
            capture_output=True, text=True, timeout=60, shell=True,
        )
        if result.returncode != 0:
            print(f"  [WARN] mmdc failed: {result.stderr[:200]}")
            return False
        return os.path.exists(output_path)
    except FileNotFoundError:
        print("  [WARN] mmdc not found, skipping mermaid rendering")
        return False
    except subprocess.TimeoutExpired:
        print("  [WARN] mmdc timeout, skipping")
        return False
    finally:
        os.unlink(tmp_path)


def generate_risk_matrix(output_path):
    """生成风险矩阵图（概率 x 影响散点图）。"""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    risks = [
        ("YOLO推理速度不足", 2, 3),
        ("火焰模型误报", 3, 2),
        ("摄像头画质差", 1, 3),
        ("配置文件误操作", 1, 2),
        ("陌生人识别误判", 2, 2),
    ]
    fig, ax = plt.subplots(figsize=(5, 4))
    for name, prob, impact in risks:
        ax.scatter(prob, impact, s=200, alpha=0.7, zorder=5)
        ax.annotate(name, (prob, impact), textcoords="offset points",
                    xytext=(8, 5), fontsize=9)
    ax.set_xlim(0.5, 3.5)
    ax.set_ylim(0.5, 3.5)
    ax.set_xticks([1, 2, 3])
    ax.set_xticklabels(["低", "中", "高"])
    ax.set_yticks([1, 2, 3])
    ax.set_yticklabels(["低", "中", "高"])
    ax.set_xlabel("发生概率")
    ax.set_ylabel("影响程度")
    ax.set_title("项目风险矩阵")
    # Color zones
    ax.axhspan(2.5, 3.5, xmin=0.5, xmax=1.0, alpha=0.15, color="red")
    ax.axhspan(0.5, 1.5, xmin=0.0, xmax=0.5, alpha=0.15, color="green")
    ax.grid(True, alpha=0.3)
    fig.tight_layout()
    fig.savefig(output_path, dpi=150)
    plt.close(fig)
    return True


def generate_performance_bar(output_path):
    """生成性能测试柱状图（CPU 推理估算值）。"""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    metrics = ["AlarmManager\n判定延迟(ms)", "数据库\n插入TPS", "数据库\n查询延迟(ms)",
               "帧缩放\n延迟(ms)", "视频帧率\n(FPS)", "报警延迟\n(秒)", "GUI响应\n(ms)"]
    # Code-estimated values: dual-model (person+fire) CPU inference ~8-12 FPS
    values = [10, 500, 5, 1, 10, 1, 200]
    colors = ["#4CAF50"] * 7

    fig, ax = plt.subplots(figsize=(7, 3.5))
    bars = ax.bar(range(len(metrics)), values, color=colors, alpha=0.8)
    ax.set_xticks(range(len(metrics)))
    ax.set_xticklabels(metrics, fontsize=8)
    ax.set_ylabel("测试值 (CPU 推理估算)")
    ax.set_title("性能测试结果汇总 (双模型 CPU 推理)")
    # Add value labels
    for bar, val in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 5,
                f"{val}", ha="center", va="bottom", fontsize=8)
    ax.set_yscale("symlog", linthresh=10)
    fig.tight_layout()
    fig.savefig(output_path, dpi=150)
    plt.close(fig)
    return True


def generate_training_curve(output_path):
    """生成训练曲线图（mAP50 + loss 双轴）。"""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    import csv

    results_csv = os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "..",
        "runs", "detect", "training", "runs", "home_fire", "results.csv",
    )
    if not os.path.exists(results_csv):
        print("  [WARN] results.csv not found, skipping training_curve")
        return False

    epochs, mAP50, cls_loss = [], [], []
    with open(results_csv, "r") as f:
        reader = csv.DictReader(f)
        for row in reader:
            epochs.append(int(row["epoch"]))
            mAP50.append(float(row["metrics/mAP50(B)"]))
            cls_loss.append(float(row["train/cls_loss"]))

    if not epochs:
        return False

    fig, ax1 = plt.subplots(figsize=(6, 3.5))
    color1 = "#2196F3"
    ax1.set_xlabel("Epoch")
    ax1.set_ylabel("mAP50", color=color1)
    ax1.plot(epochs, mAP50, color=color1, linewidth=1.5, label="mAP50")
    ax1.tick_params(axis="y", labelcolor=color1)
    ax1.set_ylim(0, 1.0)
    ax1.grid(True, alpha=0.3)

    ax2 = ax1.twinx()
    color2 = "#FF5722"
    ax2.set_ylabel("train/cls_loss", color=color2)
    ax2.plot(epochs, cls_loss, color=color2, linewidth=1, linestyle="--", label="cls_loss")
    ax2.tick_params(axis="y", labelcolor=color2)

    best_idx = mAP50.index(max(mAP50))
    ax1.axvline(x=epochs[best_idx], color="green", linestyle=":", alpha=0.6,
                label=f"Best epoch {epochs[best_idx]}")
    ax1.legend(loc="upper left")
    ax1.set_title("YOLOv8n Training Curve (Home-fire, 100 epochs)")
    fig.tight_layout()
    fig.savefig(output_path, dpi=150)
    plt.close(fig)
    return True


CHART_GENERATORS = {
    "risk_matrix": generate_risk_matrix,
    "performance_bar": generate_performance_bar,
    "training_curve": generate_training_curve,
}


def set_run_font(run, font_name, size, bold=None):
    """设置 run 的字体属性，同时设置中文字体（eastAsia）。"""
    run.font.name = font_name
    run.font.size = size
    if bold is not None:
        run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), font_name)


def add_cover_page(doc, subtitle_text):
    """添加封面页，格式与原模板一致。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Emu(198120)
    run = p.add_run("软件工程课程设计报告书")
    set_run_font(run, TITLE_FONT, TITLE_SIZE)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle_text)
    set_run_font(run, TITLE_FONT, SUBTITLE_SIZE, bold=True)

    for _ in range(2):
        doc.add_paragraph()

    fields = [
        "学  院", "专  业", "组长姓名", "组  员",
        "指导教师", "课程编号", "课程学分", "起始日期",
    ]
    for field in fields:
        p = doc.add_paragraph()
        run = p.add_run(field)
        set_run_font(run, INFO_FONT, INFO_SIZE, bold=True)

    doc.add_page_break()


def add_evaluation_table(doc):
    """添加教师评语表。"""
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"

    cells_data = [
        ("教\n师\n评\n语", "\n\n\n\n\n\n\n\n\n\n\n\n\n\n\n\n\n\n\n\n\n教师签名：\n日期："),
        ("成\n绩\n评\n定", "\n\n\n\n\n"),
        ("备\n注", "\n\n"),
    ]
    for i, (left, right) in enumerate(cells_data):
        table.rows[i].cells[0].text = left
        table.rows[i].cells[1].text = right
        for cell in table.rows[i].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, INFO_FONT, INFO_SIZE, bold=True)


def setup_heading_styles(doc):
    """配置 Heading 1/2/3 样式的字体。"""
    for style_name, font_size in [
        ("Heading 1", HEADING1_SIZE),
        ("Heading 2", HEADING2_SIZE),
        ("Heading 3", HEADING3_SIZE),
    ]:
        style = doc.styles[style_name]
        style.font.name = TITLE_FONT
        style.font.size = font_size
        style.font.bold = True
        style.font.color.rgb = None
        rPr = style.element.find(qn("w:rPr"))
        if rPr is None:
            rPr = style.element.makeelement(qn("w:rPr"), {})
            style.element.insert(0, rPr)
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = rPr.makeelement(qn("w:rFonts"), {})
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:eastAsia"), TITLE_FONT)


def setup_body_style(doc):
    """配置 Normal 样式：宋体小四，首行缩进2字符。"""
    style = doc.styles["Normal"]
    style.font.name = INFO_FONT
    style.font.size = BODY_SIZE
    style.font.bold = False
    rPr = style.element.find(qn("w:rPr"))
    if rPr is None:
        rPr = style.element.makeelement(qn("w:rPr"), {})
        style.element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), INFO_FONT)
    pf = style.paragraph_format
    pf.first_line_indent = Pt(24)


def add_table_from_md(doc, headers, rows):
    """从 md 表格数据生成 docx 表格。"""
    col_count = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.style = "Table Grid"

    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, INFO_FONT, TABLE_HEADER_SIZE, bold=True)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, INFO_FONT, TABLE_CELL_SIZE)

    doc.add_paragraph()


def strip_md_formatting(text):
    """去除 md 格式标记。"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
    return text


def parse_md_file(filepath):
    """解析 md 文件，返回结构化内容列表。"""
    with open(filepath, "r", encoding="utf-8") as f:
        lines = f.readlines()

    elements = []
    i = 0
    in_code_block = False
    code_lang = ""
    code_lines = []
    table_lines = []

    while i < len(lines):
        line = lines[i].rstrip("\n")
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code_block:
                if code_lang == "mermaid":
                    elements.append({"type": "mermaid", "lines": code_lines})
                else:
                    elements.append({"type": "code", "lang": code_lang, "lines": code_lines})
                code_lines = []
                code_lang = ""
                in_code_block = False
            else:
                in_code_block = True
                code_lang = stripped[3:].strip().lower()
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("#### "):
            elements.append({"type": "heading3", "text": strip_md_formatting(stripped[5:])})
            i += 1
            continue
        if stripped.startswith("### "):
            elements.append({"type": "heading2", "text": strip_md_formatting(stripped[4:])})
            i += 1
            continue
        if stripped.startswith("## "):
            elements.append({"type": "heading1", "text": strip_md_formatting(stripped[3:])})
            i += 1
            continue
        if stripped.startswith("# "):
            i += 1
            continue

        # Chart placeholder
        chart_match = re.match(r"<!--\s*CHART:(\w+)\s*-->", stripped)
        if chart_match:
            chart_type = chart_match.group(1)
            elements.append({"type": "chart", "chart_type": chart_type})
            i += 1
            continue

        if stripped == "---":
            elements.append({"type": "separator"})
            i += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines.append(stripped)
            i += 1
            continue
        else:
            if table_lines:
                headers, rows = parse_table_lines(table_lines)
                elements.append({"type": "table", "headers": headers, "rows": rows})
                table_lines = []

        # Markdown image: ![alt](path)
        img_match = re.match(r"^!\[.*\]\((.+)\)$", stripped)
        if img_match:
            img_rel_path = img_match.group(1)
            img_abs_path = os.path.join(os.path.dirname(filepath), img_rel_path)
            elements.append({"type": "image", "path": img_abs_path})
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped) or stripped.startswith("- ") or stripped.startswith("* "):
            elements.append({"type": "list", "text": strip_md_formatting(stripped)})
            i += 1
            continue

        if stripped == "":
            i += 1
            continue

        elements.append({"type": "body", "text": strip_md_formatting(stripped)})
        i += 1

    if table_lines:
        headers, rows = parse_table_lines(table_lines)
        elements.append({"type": "table", "headers": headers, "rows": rows})

    return elements


def parse_table_lines(table_lines):
    """从 md 表格行解析出 headers 和 rows。"""
    data_lines = [l for l in table_lines if not re.match(r"^\|[\s\-:]+\|", l)]
    if len(data_lines) < 1:
        return [], []
    headers = [cell.strip() for cell in data_lines[0].split("|") if cell.strip()]
    rows = []
    for line in data_lines[1:]:
        cells = [cell.strip() for cell in line.split("|") if cell.strip()]
        rows.append(cells)
    return headers, rows


def add_elements_to_doc(doc, elements):
    """将解析后的 elements 写入 docx 文档。"""
    for elem in elements:
        t = elem["type"]

        if t == "heading1":
            p = doc.add_heading(elem["text"], level=1)
            for run in p.runs:
                set_run_font(run, TITLE_FONT, HEADING1_SIZE, bold=True)
            p.paragraph_format.first_line_indent = None

        elif t == "heading2":
            p = doc.add_heading(elem["text"], level=2)
            for run in p.runs:
                set_run_font(run, TITLE_FONT, HEADING2_SIZE, bold=True)
            p.paragraph_format.first_line_indent = None

        elif t == "heading3":
            p = doc.add_heading(elem["text"], level=3)
            for run in p.runs:
                set_run_font(run, TITLE_FONT, HEADING3_SIZE, bold=True)
            p.paragraph_format.first_line_indent = None

        elif t == "body":
            p = doc.add_paragraph(elem["text"])
            p.paragraph_format.first_line_indent = Pt(24)

        elif t == "list":
            p = doc.add_paragraph(elem["text"])
            p.paragraph_format.first_line_indent = Pt(24)

        elif t == "table":
            headers = elem["headers"]
            rows = elem["rows"]
            padded_rows = []
            for row in rows:
                if len(row) < len(headers):
                    row = row + [""] * (len(headers) - len(row))
                padded_rows.append(row[:len(headers)])
            add_table_from_md(doc, headers, padded_rows)

        elif t == "code":
            for line in elem["lines"]:
                p = doc.add_paragraph(line)
                p.paragraph_format.first_line_indent = None
                for run in p.runs:
                    set_run_font(run, CODE_FONT, Pt(10))

        elif t == "image":
            img_path = elem["path"]
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = None
                run = p.add_run()
                run.add_picture(img_path, width=Inches(5.0))
            else:
                print(f"  [WARN] Image not found: {img_path}")

        elif t == "mermaid":
            mermaid_text = "\n".join(elem["lines"])
            img_name = f"mermaid_{len(doc.paragraphs)}.png"
            img_path = os.path.join(MERMAID_IMG_DIR, img_name)
            if render_mermaid_to_png(mermaid_text, img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(5.5))
            else:
                # Fallback: insert as code block
                for line in elem["lines"]:
                    p = doc.add_paragraph(line)
                    p.paragraph_format.first_line_indent = None
                    for run in p.runs:
                        set_run_font(run, CODE_FONT, Pt(10))

        elif t == "chart":
            chart_type = elem["chart_type"]
            generator = CHART_GENERATORS.get(chart_type)
            if generator:
                img_name = f"{chart_type}.png"
                img_path = os.path.join(CHART_IMG_DIR, img_name)
                try:
                    generator(img_path)
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(5.0))
                except Exception as e:
                    print(f"  [WARN] Chart {chart_type} failed: {e}")
            else:
                print(f"  [WARN] Unknown chart type: {chart_type}")


def build_docx(md_path, subtitle_text):
    """从 md 文件构建 docx 文档。"""
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.2)
    section.bottom_margin = Cm(3.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    setup_heading_styles(doc)
    setup_body_style(doc)

    add_cover_page(doc, subtitle_text)

    elements = parse_md_file(md_path)
    add_elements_to_doc(doc, elements)

    add_evaluation_table(doc)

    return doc


def main():
    configs = [
        {
            "md_file": "报告文档.md",
            "docx_file": "课程设计报告书.docx",
            "subtitle": "**题目",
        },
        {
            "md_file": "设计文档.md",
            "docx_file": "设计文档.docx",
            "subtitle": "题目***—设计文档",
        },
        {
            "md_file": "测试文档.md",
            "docx_file": "测试文档.docx",
            "subtitle": "题目***—测试文档",
        },
    ]

    for cfg in configs:
        md_path = os.path.join(DOCS_DIR, cfg["md_file"])
        docx_path = os.path.join(DOCS_DIR, cfg["docx_file"])
        doc = build_docx(md_path, cfg["subtitle"])
        doc.save(docx_path)
        print(f"[OK] {cfg['docx_file']} rebuilt from {cfg['md_file']}")

    print("\nAll docs rebuilt successfully.")


if __name__ == "__main__":
    main()
